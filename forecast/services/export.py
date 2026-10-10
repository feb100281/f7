"""
Выгрузка прогноза в штуках в Excel: листы «Параметры», «Группы», «Артикулы».

Данные (прогноз по месяцам, статистика спроса) — значениями, расчёт запаса — формулами Excel
от ячеек на листе «Параметры»: поменяли срок поставки или уровень сервиса — всё пересчиталось.
Компания целиком (сервис + продажи): закупка общая на центральный склад.
"""

from __future__ import annotations

from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from catalog.models import ItemGroup
from forecast.models import Series
from forecast.services import stock
from forecast.services.engine import params_label
from sales.dashboards.period import month_label

FONT = "Arial"
RED = "D3141C"
INK = "1F2937"
MUTED = "6B7280"
LINE = "E5E7EB"
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")
HEAD_FILL = PatternFill("solid", fgColor=INK)
CALC_HEAD_FILL = PatternFill("solid", fgColor=RED)
CALC_FILL = PatternFill("solid", fgColor="FDF2F2")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F6")
NEW_FILL = PatternFill("solid", fgColor="FDF3E1")   # новинка — проверить руками
THIN = Border(bottom=Side(style="thin", color=LINE))

F_QTY = '#,##0.0;-#,##0.0;"–"'
F_INT = '#,##0;-#,##0;"–"'
F_RUB = '#,##0" ₽";-#,##0" ₽";"–"'
F_CV = '0.00'
F_PCT = '0%'
F_DATE = 'DD.MM.YYYY'

P_LT, P_SL, P_Z, P_RARE = "Параметры!$B$4", "Параметры!$B$5", "Параметры!$B$6", "Параметры!$B$7"


def _font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def _title(ws: Worksheet, text: str, sub: str, width: int):
    ws["A1"] = text
    ws["A1"].font = _font(size=14, bold=True, color=RED)
    ws["A2"] = sub
    ws["A2"].font = _font(size=9, color=MUTED)
    ws.row_dimensions[1].height = 22


def _header(ws: Worksheet, row: int, headers: list[tuple[str, int, bool]]):
    for col, (text, width, calc) in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = CALC_HEAD_FILL if calc else HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 58


# ---------------------------------------------------------------------------

def _params_sheet(ws: Worksheet, run, lt, service_level, hz):
    _title(ws, "FORMULA7 · прогноз продаж запчастей и точка заказа",
           f"Прогноз от {timezone.localtime(run.created):%d.%m.%Y %H:%M}, факт продаж по {run.data_end:%d.%m.%Y}. "
           f"Выгружено {timezone.localtime():%d.%m.%Y %H:%M}.", 4)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90

    ws["A3"] = "Параметры (жёлтые ячейки можно менять — листы пересчитаются)"
    ws["A3"].font = _font(bold=True)
    rows = [
        ("Срок поставки, мес.", lt, "0", True, "Средний срок поставки (со слов клиента: срочный заказ 2–4 мес., сезонный от 6)."),
        ("Уровень сервиса", service_level, F_PCT, True, "Доля случаев, когда страхового запаса хватит до прихода поставки."),
        ("z (из уровня сервиса)", "=NORMSINV(B5)", "0.00", False, "Сколько «разбросов» спроса закладываем в страховой запас: 95% → 1,64."),
        ("Редкий спрос: меньше месяцев с продажами из 12", stock.RARE_MONTHS, "0", True,
         "Для редкого спроса страховой запас не считается — кандидат «под заказ»."),
    ]
    for i, (label, value, fmt, editable, note) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = _font()
        c = ws.cell(row=i, column=2, value=value)
        c.number_format = fmt
        c.font = _font(bold=True, color="0000FF" if editable else "000000")
        if editable:
            c.fill = INPUT_FILL
        ws.cell(row=i, column=3, value=note).font = _font(color=MUTED)

    r = 9
    ws.cell(row=r, column=1, value="Как считается").font = _font(bold=True, size=11)
    steps = [
        "1. Прогноз выручки на 6 мес. вперёд (Prophet) — отдельно сервис (наряды) и реализация + розница.",
        "2. Выручка раскладывается по товарным группам — по доле группы в этом же месяце за 2 последних сезона.",
        "3. Внутри группы — по артикулам: по штукам за последние 12 мес. (вместе с гарантией), в штуки по свежей цене.",
        ("   Старые номера одного товара (779282 → 9779282, масла с теми же характеристиками) посчитаны вместе "
         "с актуальным номером: прогноз, CV и с/с — по всему семейству, в названии — «[+ старые номера]»."
         if run.families else "   Каждый номер считается отдельно (без склейки замен номеров)."),
        "4. CV — разброс продаж артикула по 12 полным месяцам, нулевые месяцы входят: чем больше, тем рванее спрос.",
        "5. Страховой запас = z × CV × средний прогноз в месяц × √(срок поставки).",
        "6. Точка заказа = прогноз продаж на срок поставки + страховой запас (вверх до целого).",
        "   Редкий спрос — без страхового запаса и с обычным округлением: прогноз 0,2 шт. не повод держать штуку.",
        "   Когда остаток на складе + в пути опускается до точки заказа — пора заказывать.",
        "7. Себестоимость — последняя по продажам; «точка заказа по с/с» — сколько денег держать в этих штуках.",
    ]
    steps.append("8. НОВИНКИ (строки выделены, у артикула — комментарий): продажи пошли за последние полгода, "
                 "а раньше почти не было. Разброс и спрос — с месяца начала продаж; меньше 3 мес. — "
                 "«новинка: мало данных», без страхового запаса. Прогноз по короткой истории — проверить руками.")
    for i, text in enumerate(steps, start=r + 1):
        ws.cell(row=i, column=1, value=text).font = _font()
    r = r + len(steps) + 2
    ws.cell(row=r, column=1, value="Спрос (по CV за 12 мес.)").font = _font(bold=True, size=11)
    legend = [
        ("стабильный", "CV до 0,5 — продаётся ровно, держать на складе"),
        ("колеблется", "CV 0,5–1 — продаётся регулярно, но неровно"),
        ("нерегулярный", "CV больше 1 — продаётся рывками, запас большой — проверить, нужен ли"),
        ("редкий", f"продажи меньше чем в {stock.RARE_MONTHS} мес. из 12 — без страхового запаса, обычное округление, кандидат «под заказ»"),
    ]
    for i, (name, text) in enumerate(legend, start=r + 1):
        ws.cell(row=i, column=1, value=name).font = _font(bold=True)
        ws.cell(row=i, column=3, value=text).font = _font(color=MUTED)
    r = r + len(legend) + 2
    ws.cell(row=r, column=1, value="Модели прогноза выручки").font = _font(bold=True, size=11)
    for i, (s, p) in enumerate(run.params.items(), start=r + 1):
        ws.cell(row=i, column=1, value=Series(s).label).font = _font()
        ws.cell(row=i, column=3, value=params_label(p)).font = _font(color=MUTED)
    r = r + len(run.params) + 2
    ws.cell(row=r, column=1, value="Месяцы прогноза").font = _font(bold=True, size=11)
    text = ", ".join(month_label(m) for m in hz.full[:6])
    if hz.partial:
        text = f"{month_label(hz.partial)} (только остаток месяца), " + text
    ws.cell(row=r, column=3, value=text).font = _font(color=MUTED)


def _items_sheet(ws: Worksheet, rows, hz, groups, olds=None):
    olds = olds or {}
    months = hz.full[:6]
    _title(ws, "Артикулы", "Прогноз продаж по месяцам, статистика спроса и точка заказа. "
           "Красные колонки — расчёт по параметрам с листа «Параметры».", 0)
    # колонки
    headers = [
        ("Группа", 26, False), ("Артикул", 14, False), ("Название", 40, False), ("Техника", 14, False),
        ("Мес. с продажами из 12 (новинка — с начала продаж)", 13, False), ("Продано за 12 мес., шт.", 11, False),
        ("Среднее в мес., шт.", 10, False), ("Std в мес., шт.", 10, False), ("CV", 8, False), ("Спрос", 13, False),
        ("Первая продажа", 12, False), ("Последняя продажа", 12, False),
        ("Цена, ₽", 11, False), ("Откуда цена", 19, False), ("С/с последняя, ₽/шт.", 12, False),
        ("Дата с/с", 12, False),
        (f"{month_label(hz.partial)}, остаток" if hz.partial else "Остаток месяца", 10, False),
    ] + [(f"Прогноз {month_label(m)}", 10, False) for m in months] + [
        ("Прогноз продаж 6 мес., шт.", 11, False), ("в т.ч. сервис", 10, False), ("в т.ч. реализация и розница", 11, False),
        ("Прогноз на срок поставки, шт.", 11, True), ("Страховой запас, шт.", 11, True),
        ("Точка заказа, шт.", 10, True), ("Точка заказа по с/с, ₽", 14, True),
        ("Прогноз выручки 6 мес., ₽", 14, False),
    ]
    HEAD, FIRST = 4, 5
    _header(ws, HEAD, headers)
    col = {name: i for i, (name, _, _) in enumerate(headers, start=1)}
    c_m1 = 18  # первая колонка месяцев (R)
    c_m6 = c_m1 + len(months) - 1
    L = get_column_letter
    m1, m6 = L(c_m1), L(c_m6)
    # строка с номером месяца над месяцами — для «срок поставки» в SUMPRODUCT
    for i, _ in enumerate(months, start=1):
        c = ws.cell(row=HEAD - 1, column=c_m1 + i - 1, value=i)
        c.font = _font(size=8, color=MUTED)
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=HEAD - 1, column=c_m1 - 1, value="мес. №").font = _font(size=8, color=MUTED)

    c_total = c_m6 + 1
    c_lt, c_ss, c_rop, c_ropc, c_rev = c_total + 3, c_total + 4, c_total + 5, c_total + 6, c_total + 7

    for n, r in enumerate(rows):
        i = FIRST + n
        st, item = r["stats"], r["item"]
        values = [
            groups.get(r["group_id"], "Без группы"), item.article,
            f"{item.name}  [+ старые номера: {', '.join(olds[r['item_id']])}]" if olds.get(r["item_id"]) else item.name,
            item.platform.name if item.platform_id else "",
            r["months12"], st.qty12 if st else 0, st.mean12 if st else 0, st.std12 if st else 0,
            f'=IF(G{i}>0,H{i}/G{i},"")',
            # новинка: класс считан по месяцам с начала продаж — значением, не формулой
            r["demand"] if r["is_new"] else
            f'=IF(E{i}=0,"нет продаж",IF(E{i}<{P_RARE},"редкий",IF(I{i}<=0.5,"стабильный",IF(I{i}<=1,"колеблется","нерегулярный"))))',
            st.first_sale if st else None, st.last_sale if st else None,
            r["price"], r["price_label"], r["cost"], st.last_cost_date if st else None,
            r["rest"], *r["months"],
            f"=SUM({m1}{i}:{m6}{i})", r["service6"], r["shop6"],
            f"=SUMPRODUCT(({m1}${HEAD - 1}:{m6}${HEAD - 1}<={P_LT})*{m1}{i}:{m6}{i})",
            f'=IF(OR(J{i}="редкий",J{i}="нет продаж",J{i}="{stock.NEW_FEW}",I{i}=""),0,{P_Z}*I{i}*AVERAGE({m1}{i}:{m6}{i})*SQRT({P_LT}))',
            f'=IF(OR(J{i}="редкий",J{i}="нет продаж",J{i}="{stock.NEW_FEW}"),ROUND({L(c_lt)}{i},0),ROUNDUP({L(c_lt)}{i}+{L(c_ss)}{i},0))',
            f'=IF(O{i}="",0,{L(c_rop)}{i}*O{i})',
            r["revenue6"],
        ]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = _font()
            c.border = THIN
        for j in (5,):
            ws.cell(row=i, column=j).number_format = F_INT
        for j in (6, 7, 8, 17, *range(c_m1, c_m6 + 1), c_total, c_total + 1, c_total + 2, c_lt, c_ss):
            ws.cell(row=i, column=j).number_format = F_QTY
        ws.cell(row=i, column=9).number_format = F_CV
        for j in (11, 12, 16):
            ws.cell(row=i, column=j).number_format = F_DATE
        for j in (13, 15, c_ropc, c_rev):
            ws.cell(row=i, column=j).number_format = F_RUB
        ws.cell(row=i, column=c_rop).number_format = F_INT
        ws.cell(row=i, column=c_rop).font = _font(bold=True)
        ws.cell(row=i, column=c_total).font = _font(bold=True)
        for j in (c_lt, c_ss, c_rop, c_ropc):
            ws.cell(row=i, column=j).fill = CALC_FILL
        if r["is_new"]:
            for j in range(1, c_rev + 1):
                if j not in (c_lt, c_ss, c_rop, c_ropc):
                    ws.cell(row=i, column=j).fill = NEW_FILL
            note = (f"НОВИНКА: продажи пошли {r['new_since']:%d.%m.%Y}, история {r['months_active']} мес. "
                    f"Среднее, разброс (CV) и спрос — с этого месяца. Прогноз по короткой истории: "
                    f"проверьте вручную перед заказом.")
            if r["demand"] == stock.NEW_FEW:
                note += " Меньше 3 мес. продаж — страховой запас не считаем."
            c = ws.cell(row=i, column=2)
            c.comment = Comment(note, "F7")
            c.comment.width, c.comment.height = 320, 110
            c.font = _font(bold=True, color="8A5A0B")

    last = FIRST + len(rows) - 1
    # итог
    tr = last + 2
    ws.cell(row=tr, column=1, value="Итого").font = _font(bold=True)
    for j in (6, *range(c_m1 - 1, c_m6 + 1), c_total, c_total + 1, c_total + 2, c_lt, c_ss, c_rop, c_ropc, c_rev):
        c = ws.cell(row=tr, column=j, value=f"=SUBTOTAL(9,{L(j)}{FIRST}:{L(j)}{last})")
        c.font = _font(bold=True)
        c.number_format = F_RUB if j in (c_ropc, c_rev) else F_INT
    for j in range(1, c_rev + 1):
        ws.cell(row=tr, column=j).fill = TOTAL_FILL

    ws.freeze_panes = ws.cell(row=FIRST, column=4)
    ws.auto_filter.ref = f"A{HEAD}:{L(c_rev)}{last}"
    ws.conditional_formatting.add(
        f"I{FIRST}:I{last}",
        ColorScaleRule(start_type="num", start_value=0.3, start_color="E8F5E9",
                       mid_type="num", mid_value=1, mid_color="FFF8E1",
                       end_type="num", end_value=3, end_color="FDE2E2"))
    ws.conditional_formatting.add(
        f"J{FIRST}:J{last}", CellIsRule(operator="equal", formula=['"редкий"'], font=Font(name=FONT, color="9CA3AF")))
    ws.conditional_formatting.add(
        f"J{FIRST}:J{last}", CellIsRule(operator="equal", formula=[f'"{stock.NEW_FEW}"'],
                                         font=Font(name=FONT, color="8A5A0B", bold=True)))
    return {"first": FIRST, "last": last, "col": {
        "months12": "E", "qty12": "F", "demand": "J", "total": L(c_total), "lt": L(c_lt), "ss": L(c_ss),
        "rop": L(c_rop), "ropc": L(c_ropc), "rev": L(c_rev)}}


def _groups_sheet(ws: Worksheet, group_names: list[str], ref: dict):
    _title(ws, "Группы", "Суммы по артикулам с листа «Артикулы» (формулы): меняются вместе с параметрами.", 0)
    headers = [
        ("Группа", 34, False), ("Артикулов", 10, False), ("из них регулярных", 11, False),
        ("Продано за 12 мес., шт.", 12, False), ("Прогноз продаж 6 мес., шт.", 12, False),
        ("Прогноз на срок поставки, шт.", 12, True), ("Страховой запас, шт.", 12, True),
        ("Точка заказа, шт.", 12, True), ("Точка заказа по с/с, ₽", 15, True),
        ("Прогноз выручки 6 мес., ₽", 15, False), ("Доля выручки", 9, False),
    ]
    HEAD, FIRST = 4, 5
    _header(ws, HEAD, headers)
    f, l, c = ref["first"], ref["last"], ref["col"]

    def rng(letter):
        return f"Артикулы!${letter}${f}:${letter}${l}"

    last = FIRST + len(group_names) - 1
    tr = last + 1
    for n, name in enumerate(group_names):
        i = FIRST + n
        regular = "+".join(f'COUNTIFS({rng("A")},A{i},{rng(c["demand"])},"{d}")'
                           for d in ("стабильный", "колеблется", "нерегулярный"))
        values = [
            name,
            f"=COUNTIF({rng('A')},A{i})",
            f"={regular}",
            f"=SUMIFS({rng(c['qty12'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['total'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['lt'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['ss'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['rop'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['ropc'])},{rng('A')},A{i})",
            f"=SUMIFS({rng(c['rev'])},{rng('A')},A{i})",
            f"=IF($J${tr}>0,J{i}/$J${tr},0)",
        ]
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = _font(bold=(j == 8))
            cell.border = THIN
            cell.number_format = {9: F_RUB, 10: F_RUB, 11: F_PCT}.get(j, F_INT)
            if 6 <= j <= 9:
                cell.fill = CALC_FILL
    ws.cell(row=tr, column=1, value="Итого").font = _font(bold=True)
    for j in range(2, 11):
        L = get_column_letter(j)
        cell = ws.cell(row=tr, column=j, value=f"=SUM({L}{FIRST}:{L}{last})")
        cell.font = _font(bold=True)
        cell.number_format = F_RUB if j in (9, 10) else F_INT
    for j in range(1, 12):
        ws.cell(row=tr, column=j).fill = TOTAL_FILL
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    ws.auto_filter.ref = f"A{HEAD}:K{last}"


def build_workbook(run, lt: int = stock.LEAD_TIME, service_level: float = stock.SERVICE_LEVEL):
    rows, hz = stock.item_rows(run, stock.series_names(Series.TOTAL), lt, service_level)
    groups = {g.pk: g.name for g in ItemGroup.objects.all()}
    # порядок: группы по прогнозной выручке, внутри — артикулы по выручке
    rev_by_group: dict = {}
    for r in rows:
        rev_by_group[r["group_id"]] = rev_by_group.get(r["group_id"], 0) + r["revenue6"]
    rows.sort(key=lambda r: (-rev_by_group[r["group_id"]], -r["revenue6"]))
    group_names = [groups.get(g, "Без группы") for g in sorted(rev_by_group, key=lambda g: -rev_by_group[g])]

    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "Параметры"
    ws_g = wb.create_sheet("Группы")
    ws_i = wb.create_sheet("Артикулы")
    _params_sheet(ws_p, run, lt, service_level, hz)
    olds = stock.family_olds([r["item_id"] for r in rows]) if run.families else {}
    ref = _items_sheet(ws_i, rows, hz, groups, olds)
    _groups_sheet(ws_g, group_names, ref)
    for ws in (ws_p, ws_g, ws_i):
        ws.sheet_view.showGridLines = False
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    filename = f"F7_forecast_{run.data_end:%Y-%m-%d}_LT{lt}_{datetime.now():%H%M}.xlsx"
    return wb, filename
